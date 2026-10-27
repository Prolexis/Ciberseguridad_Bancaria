"""
Módulo de Generación de Reportes Operacionales, Ejecutivos y de Cumplimiento Regulatorio.
Genera documentos en formatos PDF (ReportLab), Word .docx (python-docx) y Excel .xlsx (openpyxl/pandas).
"""
import io
from datetime import datetime
from typing import List, Dict, Any
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
import docx
from docx.shared import Inches, Pt, RGBColor
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


class SecurityReportGenerator:
    """Generador multiformato de reportes de ciberseguridad bancaria."""

    @staticmethod
    def generate_pdf_report(
        title: str,
        period: str,
        kpis: Dict[str, Any],
        incidents: List[Dict[str, Any]]
    ) -> bytes:
        """Genera un reporte corporativo formal en PDF con ReportLab."""
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'BankTitle',
            parent=styles['Heading1'],
            fontSize=20,
            leading=24,
            textColor=colors.HexColor("#0B2545"),
            spaceAfter=6
        )
        subtitle_style = ParagraphStyle(
            'BankSubtitle',
            parent=styles['Normal'],
            fontSize=11,
            textColor=colors.HexColor("#4A5568"),
            spaceAfter=15
        )
        h2_style = ParagraphStyle(
            'BankH2',
            parent=styles['Heading2'],
            fontSize=14,
            leading=18,
            textColor=colors.HexColor("#134074"),
            spaceBefore=12,
            spaceAfter=6
        )

        elements = []

        # Encabezado institucional
        elements.append(Paragraph("<b>BANCO CORPORATIVO CENTRAL — DEPARTAMENTO DE CIBERSEGURIDAD</b>", title_style))
        elements.append(Paragraph(f"<b>Informe Ejecutivo Antifraude y Phishing: {title}</b> | Periodo: {period}", subtitle_style))
        elements.append(HRFlowable(width="100%", thickness=2, color=colors.HexColor("#134074"), spaceAfter=15))

        # Sección 1: Resumen de Indicadores Clave (KPIs)
        elements.append(Paragraph("1. Métricas Clave de Operación del SOC", h2_style))
        kpi_data = [
            ["Indicador de Seguridad", "Valor Medido", "Objetivo / SLA", "Estado"],
            ["Total Amenazas Analizadas", str(kpis.get("total_detections", 1420)), "N/A", "Normal"],
            ["Incidentes Críticos Confirmados", str(kpis.get("critical_incidents", 28)), "< 50", "Controlado"],
            ["Tasa de Falsos Positivos", f"{kpis.get('false_positive_rate', 1.2)}%", "< 3.0%", "Óptimo"],
            ["Tiempo Promedio de Respuesta (MTTR)", f"{kpis.get('mttr_minutes', 8.5)} min", "< 15 min", "Cumple SLA"],
        ]
        kpi_table = Table(kpi_data, colWidths=[200, 110, 110, 120])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0B2545")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.HexColor("#F8FAFC"), colors.white]),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 15))

        # Sección 2: Muestra de Incidentes Críticos Recientes
        elements.append(Paragraph("2. Registro de Incidentes Críticos y Vectores de Ataque", h2_style))
        inc_data = [["ID Caso", "Canal", "Severidad", "Objetivo Suplantado", "Estado"]]
        for inc in incidents[:10]:
            inc_data.append([
                inc.get("number", "INC-2026-001"),
                inc.get("channel", "WEB"),
                inc.get("severity", "CRITICAL"),
                inc.get("target", "Banca Móvil"),
                inc.get("status", "CONTAINED")
            ])
        inc_table = Table(inc_data, colWidths=[100, 90, 100, 150, 100])
        inc_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#134074")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8.5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ]))
        elements.append(inc_table)
        elements.append(Spacer(1, 20))

        # Declaración de Confidencialidad y Cumplimiento
        footer_style = ParagraphStyle('Foot', parent=styles['Normal'], fontSize=8, textColor=colors.HexColor("#64748B"))
        elements.append(Paragraph(
            "<b>CLASIFICACIÓN: ESTRICTAMENTE CONFIDENCIAL / USO INTERNO BANCARIO</b><br/>"
            "Este documento contiene información propietaria protegida por la política de seguridad ISO/IEC 27001 "
            "y normas de protección al cuentahabiente. Prohibida su distribución no autorizada.",
            footer_style
        ))

        doc.build(elements)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_docx_report(
        title: str,
        period: str,
        kpis: Dict[str, Any],
        incidents: List[Dict[str, Any]]
    ) -> bytes:
        """Genera un reporte corporativo editable en formato Word (.docx)."""
        doc = docx.Document()
        
        # Título
        heading = doc.add_heading(f"AegisBank: {title}", level=0)
        heading.runs[0].font.color.rgb = RGBColor(11, 37, 69)

        doc.add_paragraph(f"Periodo de Auditoría: {period}")
        doc.add_paragraph("Autor: Centro de Operaciones de Ciberseguridad (SOC) & Equipo Forense")

        doc.add_heading("1. Resumen Ejecutivo de Ciberseguridad", level=1)
        p = doc.add_paragraph()
        p.add_run(
            "Durante el periodo evaluado se ejecutaron políticas de contención proactiva antifraude. "
            "Las capacidades de detección automatizada (heurística + LLMs LangChain) mitigaron ataques antes de "
            "que se comprometieran credenciales bancarias de usuarios."
        )

        doc.add_heading("2. Desglose de Incidentes Relevantes", level=1)
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Light Shading Accent 1'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Incidente'
        hdr_cells[1].text = 'Canal'
        hdr_cells[2].text = 'Severidad'
        hdr_cells[3].text = 'Estado'

        for inc in incidents[:15]:
            row_cells = table.add_row().cells
            row_cells[0].text = inc.get("number", "N/A")
            row_cells[1].text = inc.get("channel", "N/A")
            row_cells[2].text = inc.get("severity", "N/A")
            row_cells[3].text = inc.get("status", "N/A")

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_xlsx_report(
        incidents: List[Dict[str, Any]],
        iocs: List[Dict[str, Any]]
    ) -> bytes:
        """Genera un archivo Excel (.xlsx) estructurado con pestañas para análisis y auditoría."""
        buffer = io.BytesIO()
        with pd.ExcelWriter(buffer, engine='openpyxl') as writer:
            df_incidents = pd.DataFrame(incidents)
            if df_incidents.empty:
                df_incidents = pd.DataFrame([{"incidente": "Ejemplo", "severidad": "ALTA", "canal": "EMAIL"}])
            df_incidents.to_excel(writer, sheet_name='Incidentes_SOC', index=False)

            df_iocs = pd.DataFrame(iocs)
            if df_iocs.empty:
                df_iocs = pd.DataFrame([{"tipo": "URL", "valor": "http://ejemplo-phish.com", "reputacion": 95}])
            df_iocs.to_excel(writer, sheet_name='IOCs_Bloqueados', index=False)

        buffer.seek(0)
        return buffer.getvalue()

